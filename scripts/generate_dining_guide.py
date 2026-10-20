"""Generate a one-page guest planning PDF and editable DOCX from shared copy."""
from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor
import fitz
from PIL import Image
from reportlab.lib.colors import HexColor
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph
from urllib.parse import quote

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "public/downloads/hotel-augusto-dining-guide.pdf"
DOCX = ROOT / "docs/hotel-augusto-dining-guide.docx"
EVIDENCE = ROOT / "evidence/apple-dominant-revision/dining-guide"
LOGO = ROOT / "public/brand/logo.jpeg"
PHOTO = ROOT / "public/images/dining-room-wide.webp"
INK, MUTED, BLUE, PALE = "#24262B", "#575D67", "#0066CC", "#F3F5F7"
TITLE = "Dining at Hotel Augusto"
INTRO = "A vegetarian meal at the hotel. A quiet meal in your room. Plan what works for your day in Varanasi."
SECTIONS = [
    ("At the restaurant", "Vegetarian dining is available on site. Ask the team for the current menu, prices and serving times."),
    ("In your room", "Room service is another option. Ask what can be ordered, when it is available and any applicable charges."),
]
CHECKLIST = [
    "Share your preferred date, time and the number of people in your party.",
    "Check which meals your room plan includes, especially breakfast.",
    "Tell the hotel about dietary requirements before you order.",
]
PHONE = "+91 80817 09146"
ADDRESS = "D 36/1, August Kunda Road, Godowlia, Jangambari, Varanasi 221001"
NOTE = "The hotel will confirm your dining arrangements, menu, serving times and charges."
WHATSAPP = "https://wa.me/918081709146?text=" + quote("Hello Hotel Augusto, I would like to enquire about dining. Could you share the current menu, serving times and meal inclusions?")
SOURCES = [
    ("Google hotel profile", "https://www.google.co.in/travel/hotels/entity/ChkI9IOdtNaKmJZrGg0vZy8xMWI2ZGc2bWszEAE"),
    ("MakeMyTrip property listing", "https://www.makemytrip.com/hotels/hotel_augusto-details-varanasi.html"),
]


def paragraph(c, text, x, top, width, size=10.5, color=INK, bold=False, leading=None):
    style = ParagraphStyle("copy", fontName="AugustoHeading" if bold else "AugustoBody",
                           fontSize=size, leading=leading or size * 1.45, textColor=HexColor(color))
    item = Paragraph(text, style)
    _, height = item.wrap(width, 1000)
    item.drawOn(c, x, top - height)
    return height


def pdf():
    pdfmetrics.registerFont(TTFont("AugustoBody", "C:/Windows/Fonts/arial.ttf"))
    pdfmetrics.registerFont(TTFont("AugustoHeading", str(ROOT / "public/fonts/augusto-og-inter-600-v1.ttf")))
    c = canvas.Canvas(str(PDF), pagesize=(595.276, 841.89), pageCompression=1)
    c.setTitle("Dining Planning Guide | Hotel Augusto, Varanasi")
    c.setAuthor("Digiocular for Hotel Augusto")
    c.setSubject("Verified vegetarian dining and room-service planning information")
    c.setFillColor(HexColor("#FFFFFF"))
    c.rect(0, 0, 595.276, 841.89, fill=1, stroke=0)
    c.drawImage(str(LOGO), 461, 726, width=92, height=92, preserveAspectRatio=True, anchor="c")
    paragraph(c, "HOTEL AUGUSTO · VARANASI", 42, 791, 400, 10, BLUE, True)
    paragraph(c, "Dining, around<br/>your day.", 42, 765, 390, 32, leading=35, bold=True)
    paragraph(c, INTRO, 42, 675, 490, 11, MUTED)
    with Image.open(PHOTO) as image:
        width, height = image.size
        scale = max(511 / width, 170 / height)
        photo_buffer = BytesIO()
        image.convert("RGB").save(photo_buffer, format="JPEG", quality=85, optimize=True)
        photo_buffer.seek(0)
        c.saveState()
        clip = c.beginPath()
        clip.roundRect(42, 449, 511, 170, 12)
        c.clipPath(clip, stroke=0)
        c.drawImage(ImageReader(photo_buffer), 42 + (511 - width * scale) / 2,
                    449 + (170 - height * scale) / 2, width=width * scale, height=height * scale)
        c.restoreState()
    paragraph(c, "The restaurant at Hotel Augusto.", 42, 438, 511, 8.5, MUTED)
    for index, (heading, text) in enumerate(SECTIONS):
        x = 42 + index * 268
        paragraph(c, heading, x, 400, 242, 15, bold=True)
        paragraph(c, text, x, 374, 242, 10.5, MUTED)
    c.setStrokeColor(HexColor("#D5D9DF"))
    c.line(42, 303, 553, 303)
    paragraph(c, "Before you enquire", 42, 284, 511, 17, bold=True)
    top = 252
    for index, text in enumerate(CHECKLIST, 1):
        paragraph(c, f"{index}.", 42, top, 20, 10.5, BLUE, True)
        top -= paragraph(c, text, 64, top, 489, 10.5, MUTED) + 10
    c.setFillColor(HexColor(PALE))
    c.roundRect(42, 80, 511, 86, 12, fill=1, stroke=0)
    paragraph(c, "Talk to the hotel", 58, 148, 275, 14, bold=True)
    paragraph(c, PHONE, 58, 123, 280, 14, BLUE, True)
    c.linkURL("tel:+918081709146", (58, 100, 240, 123), relative=0, thickness=0)
    paragraph(c, "Enquire on WhatsApp →", 332, 135, 200, 10.5, BLUE, True)
    c.linkURL(WHATSAPP, (330, 112, 535, 140), relative=0, thickness=0)
    paragraph(c, NOTE, 58, 101, 475, 8, MUTED)
    paragraph(c, ADDRESS + " · Opposite State Bank of India", 42, 65, 511, 8, MUTED)
    paragraph(c, "Planning information checked 20 October 2026.", 42, 45, 275, 7.5, MUTED)
    c.setFont("AugustoBody", 7.5)
    c.setFillColor(HexColor(BLUE))
    c.drawString(421, 36, "Sources: Google")
    c.drawString(489, 36, "· MakeMyTrip")
    c.linkURL(SOURCES[0][1], (450, 32, 485, 43), relative=0, thickness=0)
    c.linkURL(SOURCES[1][1], (493, 32, 553, 43), relative=0, thickness=0)
    c.showPage()
    c.save()


def link(paragraph, label, target):
    element = OxmlElement("w:hyperlink")
    element.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run, properties, color, text = OxmlElement("w:r"), OxmlElement("w:rPr"), OxmlElement("w:color"), OxmlElement("w:t")
    color.set(qn("w:val"), BLUE[1:])
    properties.append(color)
    run.append(properties)
    text.text = label
    run.append(text)
    element.append(run)
    paragraph._p.append(element)


def word():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.5)
    section.left_margin = section.right_margin = Cm(1.7)
    for name in ("Normal", "Title", "Heading 1", "List Bullet"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string(INK[1:])
    normal = document.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(6)
    document.styles["Title"].font.size = Pt(26)
    document.styles["Title"].font.bold = True
    document.styles["Heading 1"].font.size = Pt(13)
    document.styles["Heading 1"].paragraph_format.space_before = Pt(10)
    document.styles["Heading 1"].paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.text = "HOTEL AUGUSTO · VARANASI"
    header.runs[0].font.size = Pt(8)
    document.add_picture(str(LOGO), width=Inches(.9))
    document.add_paragraph(TITLE, "Title")
    document.add_paragraph(INTRO)
    buffer = BytesIO()
    with Image.open(PHOTO) as image:
        image_ratio = image.width / image.height
        image.convert("RGB").save(buffer, format="JPEG", quality=90)
    buffer.seek(0)
    picture = document.add_picture(buffer, width=Cm(17.6), height=Cm(5.3))
    crop = OxmlElement("a:srcRect")
    frame_ratio = 17.6 / 5.3
    if image_ratio < frame_ratio:
        fraction = round((1 - image_ratio / frame_ratio) * 50000)
        crop.set("t", str(fraction))
        crop.set("b", str(fraction))
    else:
        fraction = round((1 - frame_ratio / image_ratio) * 50000)
        crop.set("l", str(fraction))
        crop.set("r", str(fraction))
    picture._inline.graphic.graphicData.pic.blipFill.insert(1, crop)
    for heading, text in SECTIONS:
        document.add_paragraph(heading, "Heading 1")
        document.add_paragraph(text)
    document.add_paragraph("Before you enquire", "Heading 1")
    for text in CHECKLIST:
        document.add_paragraph(text, "List Bullet")
    document.add_paragraph("Talk to the hotel", "Heading 1")
    p = document.add_paragraph()
    link(p, PHONE, "tel:+918081709146")
    p.add_run("  ·  ")
    link(p, "Enquire on WhatsApp", WHATSAPP)
    document.add_paragraph(ADDRESS + " · Opposite State Bank of India")
    document.add_paragraph(NOTE)
    p = document.add_paragraph("Planning information checked 20 October 2026. Sources: ")
    for i, (label, target) in enumerate(SOURCES):
        if i:
            p.add_run(" · ")
        link(p, label, target)
    for run in p.runs:
        run.font.size = Pt(8)
    document.core_properties.title = TITLE
    document.core_properties.author = "Digiocular for Hotel Augusto"
    document.save(DOCX)


def main():
    for path in (PDF.parent, DOCX.parent, EVIDENCE):
        path.mkdir(parents=True, exist_ok=True)
    assert LOGO.exists() and PHOTO.exists()
    pdf()
    word()
    with ZipFile(DOCX) as archive:
        assert archive.testzip() is None
    reopened = Document(DOCX)
    assert any(p.text == TITLE for p in reopened.paragraphs)
    with fitz.open(PDF) as document:
        assert len(document) == 1
        page = document[0]
        assert PHONE in page.get_text()
        assert not [b for b in page.get_text("blocks") if b[0] < -1 or b[1] < -1 or b[2] > page.rect.width + 1 or b[3] > page.rect.height + 1]
        page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(EVIDENCE / "dining-guide-pdf.png")
    result = {
        "generated": datetime.now(timezone.utc).isoformat(),
        "pdf": PDF.relative_to(ROOT).as_posix(),
        "docx": DOCX.relative_to(ROOT).as_posix(),
        "pdfPages": 1,
        "pdfBoundsPass": True,
        "docxReopened": True,
        "visualReview": "pending",
        "sourceEvidence": "research/apple-dominant-revision/conversion-evidence/findings.json",
        "brand": "User-supplied logo.jpeg, copied to public/brand/logo.jpeg",
        "photo": "public/images/dining-room-wide.webp; provenance in research/assets/asset-manifest.json",
    }
    (EVIDENCE / "verification.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
