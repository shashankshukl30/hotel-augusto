"""Generate the editable Word twin from the research manuscript."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(__file__).resolve().parent
doc=Document()
section=doc.sections[0]
section.top_margin=section.bottom_margin=Inches(.7)
section.left_margin=section.right_margin=Inches(.75)
normal=doc.styles["Normal"]
normal.font.name="Arial"
normal.font.size=Pt(10)
normal.font.color.rgb=RGBColor.from_string("203B36")
normal.paragraph_format.space_after=Pt(6)
for name in ["Title","Heading 1","Heading 2"]:
    doc.styles[name].font.name="Arial"
    doc.styles[name].font.color.rgb=RGBColor.from_string("203B36")
for line in (ROOT/"direction.md").read_text(encoding="utf-8").splitlines():
    if not line: continue
    line=re.sub(r"\[([^]]+)\]\(([^)]+)\)",r"\1 (\2)",line)
    if line.startswith("### "): doc.add_heading(line[4:],level=2)
    elif line.startswith("## "): doc.add_heading(line[3:],level=1)
    elif line.startswith("# "): doc.add_heading(line[2:],level=0)
    elif line.startswith("- "): doc.add_paragraph(line[2:],style="List Bullet")
    else: doc.add_paragraph(line)
doc.add_page_break()
doc.add_heading("Desktop crop study",level=1)
doc.add_picture(str(ROOT/"hero-crop-1440.png"),width=Inches(6.75))
doc.add_paragraph("Research composition only. Actual Varanasi photograph by Martijn Vonk / Unsplash. Production copy and implementation remain in the application.")
doc.save(ROOT/"direction.docx")
print("Wrote direction.docx")
