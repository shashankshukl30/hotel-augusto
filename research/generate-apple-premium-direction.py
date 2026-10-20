"""Generate matching design-contract Markdown and DOCX with python-docx."""
from pathlib import Path
import importlib.util
import json
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("handoff", ROOT / "scripts/generate_handoff.py")
handoff = importlib.util.module_from_spec(spec)
spec.loader.exec_module(handoff)
qa_path = ROOT / "evidence/qa-apple-premium-production.json"
qa = json.loads(qa_path.read_text(encoding="utf-8-sig")) if qa_path.exists() else None
verification = "Final production verification is pending. Planned evidence: evidence/apple-premium/ and evidence/qa-apple-premium-production.json. Earlier revision reports do not certify these changes."
if qa and qa.get("completed"):
    s = qa["summary"]
    verification = f"Final production report: evidence/qa-apple-premium-production.json, completed {qa['completed']}. It records {s['viewportChecks']} viewport checks, {s['axeAudits']} accessibility audits, {s['interactionChecks']} interaction checks and {s['failures']} failures. See the report for its scope and remaining findings."

blocks = [
    ("title", "Hotel Augusto"),
    ("subtitle", "Apple-inspired scale & motion · 20 October 2026"),
    ("p", "The user loved the Apple design but found it too small. Restore that foundation with larger typography, a premium interface and visible motion. This supersedes the earlier rejection of Apple styling. Keep genuine photography and all 11 content pages."),
    ("h", "Visual system"),
    ("bullet", "Local Manrope display at weight 650; Inter body. Use pearl surfaces, charcoal text, a clear muted gray and a restrained blue accent. Shared OKLCH tokens in app/globals.css are authoritative."),
    ("bullet", "Interior page titles target 80–104px on desktop and 52–64px on mobile. Body copy is 18px desktop / 17px mobile; captions and labels stay readable at 14px or larger. Keep strong hierarchy, comfortable line lengths and generous white space."),
    ("bullet", "Primary pills and enquiry fields are approximately 56px high on desktop and 50px on mobile. Photography and panels use 24–32px corners. Contact input text remains 17px on mobile to avoid a miniature form; date fields align and reflow when space is insufficient."),
    ("h", "Composition and motion"),
    ("p", "Retain the genuine Varanasi hero, labelled property inset and complete page navigation. Each page keeps a distinct opening composition. Mobile adapts cropping, spacing, image order and navigation."),
    ("p", "Use shared Framer Motion entrances and scroll movement with restrained travel and opacity changes. Booking stays immediately usable. Honour reduced motion, keyboard focus and native scrolling; no compulsory loader, decorative glass or scroll hijack."),
    ("h", "Content and booking stay grounded"),
    ("p", "Verified facts remain in content/hotel.ts and source/image ledgers remain intact. The exact MakeMyTrip dated URL produced a transport error; indexed canonical content is corroboration, not proof of live availability. Google listing/photo evidence and the three additional genuine hotel photographs retain individual provenance. Destination imagery never claims a room view."),
    ("p", "WhatsApp remains an enquiry handoff with no booking engine, payment flow, automatic message sending or stored enquiry data. The public phone, active WhatsApp account, current policies and property-photo rights remain subject to the existing owner confirmations. This visual revision does not change a production domain."),
    ("h", "Verification and handoff"),
    ("p", verification),
    ("p", "Review screenshots, crops, headline wrapping, date alignment, interactive states, reduced motion and overflow. Generate final handoff twins after release QA. Preserve sources and evidence; record retained cleanup limitations."),
    ("small", "Earlier Apple and Awwwards studies remain research history. Regenerate with python research/generate-apple-premium-direction.py."),
]

markdown = ROOT / "research/apple-premium-direction.md"
docx = ROOT / "research/apple-premium-direction.docx"
markdown.write_text(handoff.markdown(blocks), encoding="utf-8")
handoff.write_docx(blocks, docx)
document = Document(docx)
document.core_properties.title = "Hotel Augusto — Apple-inspired scale and motion"
document.core_properties.subject = "20 October 2026 design contract"
document.save(docx)
assert Document(docx).paragraphs[0].text == "Hotel Augusto"
with ZipFile(docx) as archive:
    assert archive.testzip() is None
print(f"Generated and reopened {docx}.")
