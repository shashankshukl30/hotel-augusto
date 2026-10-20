"""Generate the evidence summary and its editable Word twin from one source."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
inventory = [
    ("Hero", "Actual property video, location headline, room link, live rate/offer and dated booking bar.", "Use actual Augusto property photography and a dated WhatsApp enquiry. No live prices or booking-engine claims."),
    ("Welcome", "Name/place narrative, amenity strip, overlapping property photos and practical stay facts.", "Use Augusto's Godowlia location, genuine lobby/interior photos and independently verified services."),
    ("Rooms", "All four Vagmi categories appear in full alternating photographic spreads, each with room detail links.", "Show all three Augusto categories, real bed arrangements and multiple photographs. Do not transfer Vagmi capacity, rates or category names."),
    ("Dining", "Named Flavours restaurant, video, service narrative and a dedicated dining page.", "Show Augusto's verified vegetarian restaurant and its real dining photographs. Current menu and serving times remain hotel-confirmed."),
    ("City", "Six-stop day narrative, desktop horizontal progression and mobile touch/scroll treatment.", "Use a four-part manual itinerary: river, temple, time back at the hotel, evening aarti. Native controls, no imposed clock or included-tour claims."),
    ("Guest witness", "Source-linked guest quotations, manual/automatic controls and linked platform ratings.", "Link to Augusto's verified review profiles. Publish a quote or numerical score only with current attributable evidence."),
    ("Getting here", "Airport/rail/river approach, map, location table and address.", "Use Augusto's exact address, map and arrival-mode guidance. Do not transfer Vagmi distances or promise transfers/parking."),
    ("Good to know", "Hotel amenity groups and policy definition list.", "Use verified Augusto amenities and practical questions; obtain hotel confirmation for conflicting or variable policies."),
    ("Reserve", "Stay-date booking bar, phone/WhatsApp/email, corporate-rate form and OTA alternatives.", "Use one consistent enquiry flow with dates, rooms and all guests. Exclude corporate discounts and STAAH integration until separately established."),
    ("Header and footer", "Rooms, Flavours, City and Contact routes; explicit book action; mobile contact shortcuts; language switch.", "Keep clear Augusto page navigation and persistent booking access. Do not imply a Hindi version until translated and implemented."),
    ("Room pages", "Photo gallery, amenities, meal plans, room-prefilled dates, related rooms and external booking options.", "Provide factual room details and images, comparison, a room-prefilled enquiry and related categories. No unsupported meal plans."),
    ("Dining page", "Restaurant story, gallery, menus and delivery-platform links.", "Provide a real restaurant page. Menus, named outlet and delivery links require Augusto evidence."),
    ("City and guides", "Detailed city information, practical guides and guide-level FAQs.", "Build distinct Varanasi and practical guide routes using official destination sources and verified hotel facts."),
]
findings = {
    "date": "2026-10-20",
    "brief": "Transfer most Hotel Vagmi elements into Augusto's own hotel identity and UI/UX; copy structure, never unrelated property facts.",
    "readOnlyReferences": [
        {"name": "Hotel Vagmi", "url": "https://hotelvagmi.com/", "local": r"D:\Digiocular\Websites\Hotel Vagmi\site", "sources": ["app/(en)/page.tsx", "components/sections", "components/booking", "components/pages", "components/site/Nav.tsx", "app/globals.css"]},
        {"name": "Shree Ganesha Palace", "url": "https://sgphotels.com/", "local": r"D:\Digiocular\Websites\shree-ganesha-palace", "sources": ["app/page.tsx", "app/layout.tsx", "app/globals.css", "lib/site.ts"]},
        {"name": "The Drake", "url": "https://thedrake.ca/", "awwwards": "https://www.awwwards.com/sites/the-drake-hotel", "awardDate": "2025-09-02"},
    ],
    "homepageOrder": ["HeroSection", "WelcomeSection", "RoomsSection", "FlavourSection", "CitySection", "WitnessSection", "GettingHereSection", "GoodToKnowSection", "ReserveSection"],
    "inventory": [{"module": name, "vagmiEvidence": original, "augustoAdaptation": adaptation} for name, original, adaptation in inventory],
    "patterns": [
        "Property first: Vagmi and Shree Ganesha Palace immediately show their real hotel with a visible reservation action. Room, food and arrival utility make the experience unmistakably hospitality.",
        "Inventory in the open: Vagmi gives each room category a full photographic treatment. Guests can compare the options while scrolling instead of discovering that more options exist behind one featured card.",
        "One continuous hotel story: property introduction, rooms, dining, city, guest evidence, directions, policy and enquiry form one useful journey. The Drake reinforces this with a precise navigation taxonomy and human hospitality imagery.",
    ],
    "augustoDeficienciesAtReview": [
        "The previous large Apple-style treatment gave similar visual weight to abstract headings and selected photos; its place-first hero and limited hotel chapters did not establish the property quickly enough.",
        "A single selected-room presentation made Augusto's inventory feel smaller and less comparable than Vagmi's complete room spreads.",
        "The existing short city page and narrow practical journey left the site with less hospitality breadth despite its real separate routes.",
    ],
    "direction": "A warm editorial hotel identity: ivory and ink blue, restrained bronze details, Bodoni Moda display with Manrope body, honest property photography, full room spreads, useful arrival and stay details, small corners and generous spacing. Manual city tabs and restrained image/reveal motion provide character without hiding content.",
    "limitations": [
        "DDD Hotel's live page remained in its numeric loading/blank transition during capture. Its Awwwards listing was recorded, but it was not used as a successful live visual reference.",
        "The Scott Resort & Spa Awwwards record was read (2019-12-08); its live website was not visually inspected during this bounded study.",
        "Vagmi does not have a standalone homepage comparison matrix or homepage FAQ in the inspected source. Augusto comparison/FAQ additions should not be described as exact copies of those modules.",
        "Vagmi's property video, STAAH rates, offers, corporate terms, reviews, occupancy, restaurant identity and translation do not establish equivalent Augusto facilities or integrations.",
    ],
    "captureEvidence": "live-inspection.json",
}
(ROOT / "reference-findings.json").write_text(json.dumps(findings, indent=2, ensure_ascii=False), encoding="utf-8")

sections = [
    ("Purpose", [findings["brief"]]),
    ("Three reference patterns", findings["patterns"]),
    ("Where the previous Augusto treatment fell short", findings["augustoDeficienciesAtReview"]),
    ("Recommended direction", [findings["direction"]]),
    ("Exact Vagmi homepage order", [" → ".join(findings["homepageOrder"])]),
]
sections.append(("Element inventory and Augusto adaptation", [f"{name}. Vagmi: {original} Augusto: {adaptation}" for name, original, adaptation in inventory]))
sections.append(("Sources and inspected evidence", [
    "Hotel Vagmi — https://hotelvagmi.com/ — local read-only source: D:\\Digiocular\\Websites\\Hotel Vagmi\\site. Actual desktop and mobile captures saved; property video, dates bar, room spreads and headings verified.",
    "Shree Ganesha Palace — https://sgphotels.com/ — local read-only source: D:\\Digiocular\\Websites\\shree-ganesha-palace. Desktop hero and introductory architecture capture visually inspected. Its 1916 heritage claim belongs to that property.",
    "The Drake — https://thedrake.ca/ — https://www.awwwards.com/sites/the-drake-hotel (Site of the Day, 2 September 2025). Desktop and mobile live hero visually inspected.",
    "Capture inventory, landed URLs, headings, links and typography are retained in live-inspection.json. Browser contexts were closed after capture.",
    "Varanasi hero photograph already integrated by root: Martijn Vonk, Unsplash, https://unsplash.com/photos/a-group-of-boats-floating-on-top-of-a-river-wloNuC7qKf8. This is destination imagery, not a hotel view.",
]))
sections.append(("Coverage limits", findings["limitations"]))

title = "Hotel Augusto — Hospitality Reference Study"
markdown = [f"# {title}", "", "20 October 2026 · Digiocular", ""]
for heading, paragraphs in sections:
    markdown.extend([f"## {heading}", ""])
    for paragraph in paragraphs:
        markdown.extend([paragraph, ""])
(ROOT / "reference-study.md").write_text("\n".join(markdown), encoding="utf-8")

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.7)
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10)
style.paragraph_format.space_after = Pt(6)
style.paragraph_format.line_spacing = 1.1
for name in ["Title", "Heading 1"]:
    doc.styles[name].font.name = "Calibri"
    doc.styles[name].font.color.rgb = RGBColor.from_string("24313E")
doc.add_heading(title, 0)
doc.add_paragraph("20 October 2026 · Digiocular")
for heading, paragraphs in sections:
    doc.add_heading(heading, 1)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
footer = section.footer.paragraphs[0]
footer.text = "Hotel Augusto · Research and implementation reference · Digiocular"
footer.style = doc.styles["Caption"]
output = ROOT / "reference-study.docx"
doc.save(output)
reopened = Document(output)
assert len(reopened.paragraphs) == len(doc.paragraphs)
assert all(heading in "\n".join(p.text for p in reopened.paragraphs) for heading, _ in sections)
print(json.dumps({"json": "reference-findings.json", "markdown": "reference-study.md", "docx": output.name, "paragraphsVerified": len(reopened.paragraphs)}))
