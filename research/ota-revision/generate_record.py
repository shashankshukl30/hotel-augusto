"""Generate the matched OTA-source review record with python-docx."""
from pathlib import Path
import importlib.util
import json
from zipfile import ZipFile
from docx import Document
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
spec = importlib.util.spec_from_file_location("handoff", ROOT / "scripts/generate_handoff.py")
handoff = importlib.util.module_from_spec(spec)
spec.loader.exec_module(handoff)

mapping_path = HERE / "source-to-page-mapping.json"
mapping = json.loads(mapping_path.read_text(encoding="utf-8"))
exact = "https://www.makemytrip.com/hotels/hotel-details/?hotelId=201909301024364424&_uCurrency=INR&checkin=11192026&checkout=11202026&city=CTVNS&country=IN&lat=25.30833&lng=83.00653&locusId=CTVNS&locusType=city&rank=1&regionNearByExp=3&roomStayQualifier=2e0e&rsc=1e2e0e&searchText=HOTEL+AUGUSTO&topHtlId=201909301024364424&mtkeys=undefined&isPropSearch=T"
mapping["sources"][1]["url"] = exact
for photo in mapping["photo_to_page"]:
    if photo["google_photo"] == 3:
        photo["pages"] = ["/", "/gallery"]
        photo["match"] = "New real room photograph used in the homepage hotel inset and gallery; no room-category assignment inferred."
    elif photo["google_photo"] == 5:
        photo["pages"] = ["/gallery"]
        photo["match"] = "New real two-double-bed room photograph used in the gallery; no category assignment inferred."
    elif photo["google_photo"] == 10:
        photo["pages"] = ["/", "/dining"]
mapping_path.write_text(json.dumps(mapping, indent=2, ensure_ascii=False), encoding="utf-8")

blocks = [
    ("title", "Hotel Augusto"),
    ("subtitle", "Google & MakeMyTrip source review · 19 October 2026"),
    ("p", "The Google profile was opened successfully and its accessible photographs were inspected. MakeMyTrip's canonical indexed property content was used for factual corroboration. The exact dated link supplied by the user was attempted separately; its live offer was not recovered."),
    ("h", "What was actually accessible"),
    ("p", f"[User-supplied MakeMyTrip route]({exact}) — hotelId 201909301024364424, 19–20 November 2026. Fresh Microsoft Edge navigation returned net::ERR_HTTP2_PROTOCOL_ERROR and the browser error page, without a hotel heading. A separate HTTP/1 photo-page request timed out after 35 seconds with no bytes. This records an observed transport failure, not a proven authentication block. See ../requested-source-verification/makemytrip-exact.json and its screenshot."),
    ("p", "[Canonical MakeMyTrip property](https://www.makemytrip.com/hotels/hotel_augusto-details-varanasi.html) — a fresh search exposed indexed content crawled two weeks earlier. Matching address, Deluxe / Deluxe Twin Bed / Family categories, vegetarian restaurant, Wi-Fi, room service and lift support the website's durable facts. Indexed rates and dated terms do not establish current availability. [Photo index](https://www.makemytrip.com/hotels/photos-of-hotel_augusto-details-varanasi.html) identifies 68 property and 137 guest photos, but individual MMT image URLs were not recovered."),
    ("p", "Google Maps — the rendered heading, D 36/1 August Kunda Road/Jangambari address and +91 80817 09146 phone match the hotel. An owner section included an October 3, 2026 post. Eight still-photo entries were opened individually; the gallery's further-photo action displayed a sign-in prompt. No sign-in, message or access-control workaround was attempted. See ../requested-source-verification/google-profile.json, google-photo-candidates.json and google-gallery-expanded.png."),
    ("h", "Where the evidence is used"),
    ("bullet", "Address and telephone: /hotel, /contact and shared footer. Room names: /rooms and the three room-detail pages. Vegetarian dining: /dining and the homepage; meal inclusion and serving times remain hotel enquiries."),
    ("bullet", "Google photo 9 shows the same daytime exterior photograph retained on /hotel. Photo 7 corroborates the tan-headboard double-room photograph used in the room pages. Their existing higher-resolution originals were retained; matching is visual, not byte-level."),
    ("bullet", "New google-double-room.webp: / homepage hotel inset and /gallery; uploader Augusto Hotel, July 2026. New google-family-room.webp: /gallery; uploader HOTEL AUGUSTO VARANASI, January 2022. These photos do not establish a booking-category assignment."),
    ("bullet", "New google-dining-room.webp: / homepage dining preview and /dining detail. The Hotels.com original was linked from Google's gallery. Exact source URLs, dimensions and SHA256 hashes are in asset-candidates.json; the eight-photo comparison is google-property-contact-sheet.jpg."),
    ("h", "Limits and handoff"),
    ("p", "No MMT download provenance is claimed for photos obtained elsewhere. Public visibility and a hotel-like uploader name do not establish image reuse rights; the client must confirm permission. Conflicting policies, fixed travel times, star ratings, live prices and unverified facilities remain excluded. The public phone is verified; WhatsApp registration and hotel control still need owner confirmation. No enquiry was sent."),
    ("p", "Source-to-page-mapping.json provides the complete route/fact/photo mapping. The two new hotel/dining routes passed local browser image, overflow and runtime checks at 390px and 1440px, with four full-page screenshots visually inspected. The final dining-photo change passed both widths. See ../../evidence/hospitality-review/. Final release evidence is maintained by the main production QA report."),
]

md = HERE / "SOURCE-REVIEW.md"
docx = HERE / "SOURCE-REVIEW.docx"
md.write_text(handoff.markdown(blocks), encoding="utf-8")
handoff.write_docx(blocks, docx)
reopened = Document(docx)
reopened.styles["Normal"].font.size = Pt(10)
reopened.styles["Normal"].paragraph_format.line_spacing = 1.08
reopened.styles["Normal"].paragraph_format.space_after = Pt(5)
reopened.styles["Heading 1"].font.size = Pt(13)
reopened.styles["Heading 1"].paragraph_format.space_before = Pt(11)
reopened.core_properties.title = "Hotel Augusto — Google and MakeMyTrip source review"
reopened.core_properties.subject = "Verified source access, photo provenance and website route mapping"
reopened.save(docx)
reopened = Document(docx)
assert reopened.paragraphs[0].text == "Hotel Augusto"
with ZipFile(docx) as archive:
    assert archive.testzip() is None
print(f"Generated and reopened {docx}; {len(reopened.paragraphs)} paragraphs.")
