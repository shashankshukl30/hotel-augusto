"""Write the current design contract and its editable Word twin."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
CONTENT = r'''# Hotel Augusto: Apple-inspired design contract

Research and implementation review: 19 October 2026. This supersedes the earlier serif/terracotta direction in design-references.md. The user's explicit Apple direction governs the revision.

## Live references inspected

- [Apple iPhone](https://www.apple.com/iphone/): fresh Edge captures at 1440 and 390 CSS pixels. The desktop category title measured 80/84px, weight 600; product titles measured 56/60px. Mobile category title measured 48/52px, with product titles at 32/36px. A desktop photographic tile measured 1260px wide with 28px corners. Side margins were 6.25% at both tested widths. The important observation is the clear hierarchy around a small number of large images.
- [Apple MacBook Pro](https://www.apple.com/macbook-pro/): inspected desktop hero and mobile feature section. Section headings measured 48/52px on desktop and 32/36px on mobile. Mobile feature panels were 321.25 by 480px with 28px corners and a glimpse of the next panel. Body text measured 17/25px. This shows deliberate mobile recomposition and controlled disclosure of detail.
- [Apple HIG: Typography](https://developer.apple.com/design/human-interface-guidelines/typography): legibility, a limited typeface palette, and hierarchy through size, weight, and color. Its native-platform point sizes are not CSS specifications.
- [Apple HIG: Layout](https://developer.apple.com/design/human-interface-guidelines/layout): group related content, maintain alignment, adapt to available space, and keep controls clear of obstructed screen areas. These principles inform the website; native interface conventions are not copied literally.

Screenshots and computed browser measurements are retained in research/apple-reference/. Actual inspected files include iphone-1440-hero.png, iphone-390-hero.png, macbook-pro-1440-hero.png, macbook-pro-390-section.png, hig-layout.png, hig-typography.png, and computed-styles.json. No Apple images or product copy are included in the public hotel site.

## Augusto design values

The values below are the hotel's original adaptation, not an official Apple token system. Inter variable replaces the former display serif. Both display and body tokens point to the same family; hierarchy comes from scale, weight, spacing, and layout.

| Role | Desktop | Mobile | Weight and behavior |
| --- | --- | --- | --- |
| Page title | 72px / 1.055 | 52px / 1.06; 46px at 360 | 650; tracking -0.055em to -0.045em |
| Section title | 56px / 1.08 | 36px / 1.13 | 650; tracking -0.045em to -0.035em |
| Feature title | 40px / 1.12 | 32px / 1.13 | 650; tracking -0.035em to -0.03em |
| Intro lead | 28px / 1.28 | 23px / 1.3 | 550; restrained line length |
| Body | 17px / 1.55 | 16px / 1.55 | 400; short factual paragraphs |
| Context label | 17px / 1.35 | 14px / 1.35 | 600; sentence case |
| Caption | 12px / 1.4 | 11px / 1.4 | 400; supplementary information |

Keep shared color tokens in OKLCH: paper near-white, paper-deep cool light gray, ink charcoal, muted readable gray, and one blue accent. Use the shared theme as the authority. Avoid warm editorial ornaments, italic emphasis, widely tracked capitals, section numbering, and decorative gradients.

Use an 8px spacing rhythm, with deliberate 24/28px photo gaps. Generous sections use 104-112px desktop separation and 64-72px mobile separation. Page intros begin at 88px desktop and 60px mobile below the shared header. Main gutters stay around 6vw with a practical mobile minimum. Body copy should remain below 65-75 characters per line; compact introductory copy is shorter.

Photo and panel corners are 24-28px, with 20px gallery corners on mobile. Primary actions are blue capsules with at least 44px height. Secondary actions are plain blue text with a small arrow. Image controls are quiet 44px circles. Small decorative buttons must not replace readable booking text.

## Composition and booking behavior

- Lead with the hotel and its actual rooms. Use large genuine photographs, concise room names, and a direct booking action. Do not invent amenities, nightly prices, awards, room availability, river views, or a luxury classification.
- Keep WhatsApp booking easy to find in the shared navigation and relevant room/contact contexts. Use a prefilled inquiry with room choice where available. The label should communicate booking or availability, without implying a completed reservation.
- Gallery: one even two-column photo grid on desktop, large single-column photographs on mobile, and a compact All / Hotel / Rooms / Dining segmented control. Captions identify the real space. Native full-image dialog supports keyboard navigation, focus trapping and restoration, Escape, and horizontal touch swipes.
- Varanasi: a large river panorama, one wide dawn story, then temple and evening cards. Desktop dawn imagery and copy sit beside one another; mobile gives the image its own composition above the copy. The actual temple photo retains its attribution. Generic river photographs are never described as a specific ghat or an actual aarti ceremony.
- Keep one clear next step per major composition. Space comes from content hierarchy rather than empty ornamental sections. The hotel's practical address and route information remain visible and factual.

## Motion and interaction

Use the shared reduced-motion-aware Reveal component for restrained entrances. A suitable range is 12-18px travel over 0.5-0.8 seconds; this is a proposed hotel behavior, not a measured Apple timing. Gallery image hover scales only to 1.025 over 0.5 seconds. Controls use 0.2-second state changes and visible keyboard focus. Reduced motion removes the extra transforms. Do not add forced scroll sequences, loading theatrics, decorative parallax, or autoplaying carousels merely to suggest premium production.

## Acceptance evidence

The redesigned gallery and destination routes were checked in fresh Edge contexts at 360, 390, 768, and 1440px. Layout verification requires body scroll width to equal client width, all images to decode, and no page errors. Gallery verification covers all four filters, arrow-key and touch navigation, native-dialog focus management, and Escape return. Destination anchors river, temple, and aarti remain valid. Current screenshots and results are in research/design-qa/apple-redesign/.
'''

def clean(text: str) -> str:
    return re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)

def main() -> None:
    (ROOT / "apple-direction.md").write_text(CONTENT, encoding="utf-8")
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    table_lines = []
    for line in CONTENT.splitlines() + [""]:
        if line.startswith("|"):
            if not re.match(r"\|\s*---", line):
                table_lines.append([x.strip() for x in line.strip("|").split("|")])
            continue
        if table_lines:
            table = doc.add_table(rows=1, cols=len(table_lines[0]))
            table.style = "Light Shading Accent 1"
            for cell, text in zip(table.rows[0].cells, table_lines[0]):
                cell.text = text
            for row in table_lines[1:]:
                for cell, text in zip(table.add_row().cells, row):
                    cell.text = text
            table_lines = []
        if not line:
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("- "):
            doc.add_paragraph(clean(line[2:]), style="List Bullet")
        else:
            doc.add_paragraph(clean(line))
    doc.save(ROOT / "apple-direction.docx")
    print("Wrote research/apple-direction.md and apple-direction.docx")

if __name__ == "__main__":
    main()
