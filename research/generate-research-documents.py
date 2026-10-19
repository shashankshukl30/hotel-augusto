"""Regenerate the research handoff and its Word twins from the source JSON."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parent
FACTS = json.loads((ROOT / "hotel-facts.json").read_text(encoding="utf-8"))
LEDGER = json.loads((ROOT / "source-ledger.json").read_text(encoding="utf-8"))


def write_document(stem, title, sections):
    md = [f"# {title}", "", "Research date: 19 October 2026. Public-source evidence; not hotel management confirmation.", ""]
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    sec.left_margin = sec.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ("Title", "Heading 1", "Heading 2"):
        doc.styles[name].font.color.rgb = RGBColor.from_string("713F2B")
    doc.add_heading(title, 0)
    doc.add_paragraph(md[2])
    for heading, paragraphs in sections:
        md.extend([f"## {heading}", ""])
        doc.add_heading(heading, 1)
        for para in paragraphs:
            md.extend([para, ""])
            doc.add_paragraph(para)
    (ROOT / f"{stem}.md").write_text("\n".join(md), encoding="utf-8")
    doc.save(ROOT / f"{stem}.docx")
    with ZipFile(ROOT / f"{stem}.docx") as archive:
        assert archive.testzip() is None
        assert "word/document.xml" in archive.namelist()
    reopened = Document(ROOT / f"{stem}.docx")
    assert len(reopened.paragraphs) >= len(sections) + 2


def main():
    identity = FACTS["identity"]
    sections = [
        ("Recommended positioning", [
            "A well-located Varanasi stay in Godowlia, within walking distance of Shri Kashi Vishwanath Temple and Dashashwamedh Ghat. Lead with location, real rooms, vegetarian dining and useful arrival guidance. The quality of the website must not imply a luxury resort or riverfront property.",
            "Use the actual property photographs. Destination photographs should be clearly presented as Varanasi, never as views from the hotel. Do not publish invented guest reviews or a combined rating.",
        ]),
        ("Identity and booking contact", [
            f"Name: {identity['displayName']}. Address: {identity['address']}",
            f"Best corroborated public phone: {identity['phoneDisplay']} (S01/S03/S12). A different phone appears on Trip.com (S09); the Google-listed number is the stronger candidate.",
            "WhatsApp account registration is not verified. No message was sent. No verified email was recovered. Ask for the client's booking number before production launch if it differs from the Google-listed telephone.",
            "Legacy hotelaugusto.in did not resolve. thehotelaugusto.in redirected to a Booking.com affiliate route on 19 October 2026. Neither should be assigned as canonical solely on the basis of older listings.",
        ]),
        ("Room categories", [
            "Deluxe Room: double bed. Deluxe Twin Room: two single beds. Family Room: described by Goibibo as a four-bed family room with two king beds. MMT also lists these three categories (S02/S07).",
            "Use AC, private bathroom, television and Wi-Fi as general room/hotel facilities. Confirm exact dimensions, adult/child occupancy, extra beds and per-category amenities with the hotel; public inventory is inconsistent.",
        ]),
        ("Supported facilities", [
            "On-site restaurant, vegetarian dining, room service, 24-hour front desk, Wi-Fi, air conditioning, housekeeping and luggage assistance/storage are corroborated across listings. MMT also lists an elevator. An elevator does not establish complete step-free access.",
            "Do not promise a pool, spa, gym, river view, free parking, free breakfast on every rate, guaranteed early check-in, or fixed response times. Airport transfer availability and charges should be confirmed during enquiry.",
        ]),
        ("Policies requiring confirmation", [f"{item['topic']}: {item['evidence']} Website decision: {item['decision']}" for item in FACTS["policyConflicts"]]),
        ("Research scope and remaining unknowns", [
            "Reviewed Google Hotels, MakeMyTrip, Goibibo, Booking.com, Tripadvisor, Agoda, Trip.com, Expedia, Hotels.com, EaseMyTrip, Trivago and public directories; checked legacy website domains. Booking direct content hit robot verification, so only its indexed public excerpt was used. Agoda English returned a shell; an indexed Japanese property page supplied corroboration.",
            "Agent Reach Exa/Jina were available. Doctor found no Facebook/Instagram backend; no authentication or private account access was attempted. No verified social profile, owner email or logo master was recovered by this fact-research task. Image retrieval is tracked separately by the asset manifest.",
            "See source-ledger.json for URLs, access quality and dated evidence; hotel-facts.json contains the full structured fact/conflict matrix.",
        ]),
    ]
    write_document("hotel-fact-sheet", "Hotel Augusto | Research fact sheet", sections)
    source_sections = []
    for source in LEDGER["sources"]:
        paragraphs = [f"URL: {source['url']}", f"Source type: {source['type']}. Access: {source['access']}.", source["findings"]]
        if source.get("alternateUrl"):
            paragraphs.insert(1, f"Alternate indexed URL: {source['alternateUrl']}")
        source_sections.append((f"{source['id']} | {source['title']}", paragraphs))
    source_sections.extend([
        ("Access limits", LEDGER["notRecovered"]),
        ("Image provenance", [LEDGER["assetHandoff"]]),
        ("Tool status", [LEDGER["agentReach"]]),
    ])
    write_document("source-ledger", "Hotel Augusto | Public source ledger", source_sections)
    print("Generated and reopened fact sheet and source ledger, both Markdown and DOCX.")


if __name__ == "__main__":
    main()
